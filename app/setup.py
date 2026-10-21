"""설정 화면(첫 실행 설정·관리자 '기관 설정')이 쓰는 도우미.

.env·자산대장·공간목록·출력 양식·로고를 화면에서 바꿀 수 있게 한다.
올린 엑셀은 먼저 '.upload.xlsx'로 받아 검사하고, [저장]을 누를 때 원래 이름으로 바꾼다.
"""
from __future__ import annotations

import re
import socket
from collections import Counter
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from . import excel, locations
from .catalog import read_ledger
from .config import (ASSETS_DIR, BRANDING_DIR, LOGO_TYPES, ROOT, SOURCE_MASTER, SOURCE_SPACES,
                     SOURCE_TEMPLATE)

ENV_PATH = ROOT / ".env"
ENV_EXAMPLE = ROOT / ".env.example"
PENDING = {"ledger": (ASSETS_DIR / "고정자산대장.upload.xlsx", SOURCE_MASTER),
           "template": (ASSETS_DIR / "재물조사표.upload.xlsx", SOURCE_TEMPLATE)}
MAX_EXCEL_BYTES = 20 * 1024 * 1024
MAX_LOGO_BYTES = 2 * 1024 * 1024
SAMPLE_MARK = "SAMPLE-DUMMY"
_KEY = re.compile(r"^\s*([A-Z][A-Z0-9_]*)\s*=(.*)$")


# ── .env ──
def read_env() -> dict[str, str]:
    path = ENV_PATH if ENV_PATH.exists() else ENV_EXAMPLE
    values = {}
    if path.exists():
        for line in path.read_text(encoding="utf-8-sig").splitlines():
            match = _KEY.match(line)
            if match:
                values[match.group(1)] = match.group(2).strip()
    return values


def write_env(updates: dict[str, str]) -> bool:
    """주석과 다른 키는 그대로 두고 값만 바꾼다. 바뀐 값이 있으면 True."""
    source = ENV_PATH if ENV_PATH.exists() else ENV_EXAMPLE
    lines = source.read_text(encoding="utf-8-sig").splitlines() if source.exists() else []
    before = read_env()
    clean = {k: re.sub(r"[\r\n]", "", v).strip() for k, v in updates.items()}
    done = set()
    for i, line in enumerate(lines):
        match = _KEY.match(line)
        if match and match.group(1) in clean:
            lines[i] = f"{match.group(1)}={clean[match.group(1)]}"
            done.add(match.group(1))
    lines += [f"{k}={v}" for k, v in clean.items() if k not in done]
    ENV_PATH.write_text("\r\n".join(lines) + "\r\n", encoding="utf-8")
    return any(before.get(k, "") != v for k, v in clean.items())


def lan_url(port: str) -> str:
    """같은 와이파이의 휴대폰이 접속할 주소. UDP connect는 실제로 패킷을 보내지 않는다."""
    try:
        with socket.socket(socket.AF_INET, socket.SOCK_DGRAM) as probe:
            probe.connect(("10.255.255.255", 1))
            host = probe.getsockname()[0]
    except OSError:
        host = "localhost"
    return f"http://{host}:{port}"


# ── 엑셀 검사 ──
def is_sample(path: Path) -> bool:
    try:
        workbook = openpyxl.load_workbook(path, read_only=True)
        try:
            return workbook.properties.keywords == SAMPLE_MARK
        finally:
            workbook.close()
    except Exception:
        return False


DIFF_FIELDS = ("name", "raw_location", "book_qty", "category")


def check_ledger(path: Path, result_codes=(), current: dict | None = None) -> dict:
    """고정자산대장을 읽어 요약한다. current(지금 자산)를 주면 바뀌는 내용도 센다. 형식이 틀리면 ValueError."""
    rows = read_ledger(path)
    if not rows:
        raise ValueError("자산 행이 하나도 없습니다. '자산코드' 아래에 자산이 있는지 확인해 주세요.")
    counts = Counter(row["code"] for row in rows)
    places = Counter(row["raw_location"] for row in rows)
    unassigned = places.pop(locations.UNASSIGNED, 0)
    missing = sorted(set(result_codes) - set(counts))
    summary = {
        "assets": len(counts),
        "duplicates": sorted(code for code, n in counts.items() if n > 1)[:20],
        "locations": [{"name": name, "assets": n} for name, n in sorted(places.items())],
        "unassigned": unassigned,
        "missing_results": len(missing), "missing_examples": missing[:5],
    }
    if current is not None:
        new = {row["code"]: row for row in reversed(rows)}   # 중복이면 첫 행(카탈로그와 같은 규칙)
        added = [c for c in new if c not in current]
        removed = [c for c in current if c not in new]
        changed = [c for c in new if c in current
                   and any(new[c].get(f) != current[c].get(f) for f in DIFF_FIELDS)]
        label = lambda c, src: f"{c} {src[c]['name']}"
        summary["diff"] = {
            "added": len(added), "removed": len(removed), "changed": len(changed),
            "examples": {"added": [label(c, new) for c in added[:5]],
                         "removed": [label(c, current) for c in removed[:5]],
                         "changed": [label(c, new) for c in changed[:5]]},
        }
    return summary


def check_template(path: Path) -> dict:
    """재물조사표 양식의 머리글을 확인한다. 필수 머리글이 없으면 ValueError."""
    workbook = openpyxl.load_workbook(path)
    try:
        sheet = workbook.worksheets[0]
        columns, start = excel.find_columns(sheet)
        filled = sum(1 for r in range(start, sheet.max_row + 1) if sheet.cell(r, columns["code"]).value)
    finally:
        workbook.close()
    return {"columns": [excel.LABELS[k] for k in excel.LABELS if k in columns], "filled_rows": filled}


def save_upload(kind: str, content: bytes) -> Path:
    if len(content) > MAX_EXCEL_BYTES:
        raise ValueError("엑셀 파일은 20MB 이하여야 합니다.")
    if not content.startswith(b"PK"):
        raise ValueError(".xlsx 파일만 올릴 수 있습니다. 엑셀에서 '다른 이름으로 저장 → Excel 통합 문서'로 저장해 주세요.")
    pending = PENDING[kind][0]
    pending.write_bytes(content)
    return pending


def discard_uploads() -> None:
    for pending, _target in PENDING.values():
        pending.unlink(missing_ok=True)


def apply_uploads() -> list[str]:
    applied = []
    for kind, (pending, target) in PENDING.items():
        if pending.exists():
            pending.replace(target)
            applied.append(kind)
    return applied


# ── 로고 ──
def save_logo(filename: str, content: bytes) -> None:
    suffix = Path(filename or "").suffix.lower()
    if suffix not in LOGO_TYPES:
        raise ValueError("로고는 svg·png·webp·jpg 파일만 쓸 수 있습니다.")
    if len(content) > MAX_LOGO_BYTES:
        raise ValueError("로고 파일은 2MB 이하여야 합니다.")
    BRANDING_DIR.mkdir(exist_ok=True)
    for old in BRANDING_DIR.glob("logo.*"):
        old.unlink()
    (BRANDING_DIR / f"logo{'.jpg' if suffix == '.jpeg' else suffix}").write_bytes(content)


def remove_logo() -> None:
    for old in BRANDING_DIR.glob("logo.*"):
        old.unlink()


# ── 공간목록 ──
def write_spaces(items: list[dict]) -> None:
    """[{floor, name}] → 공간목록.xlsx (구역 순서는 처음 나온 순서를 따른다)."""
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "공간목록"
    thin = Side(style="thin", color="FFB7B7B7")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    sheet.append(["구역", "공간명", "비고"])
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="FFD9E5E9")
        cell.font = Font(name="맑은 고딕", bold=True)
        cell.alignment = Alignment(horizontal="center")
    order = list(dict.fromkeys(item["floor"] for item in items))
    for floor in order:
        for item in items:
            if item["floor"] == floor:
                sheet.append([floor, item["name"], None])
    for row in sheet.iter_rows(min_row=1):
        for cell in row:
            cell.border = border
    sheet.column_dimensions["A"].width = 14
    sheet.column_dimensions["B"].width = 24
    sheet.column_dimensions["C"].width = 30
    sheet.freeze_panes = "A2"
    workbook.save(SOURCE_SPACES)


def current_spaces() -> list[dict]:
    return [{"floor": floor, "name": name} for floor, names in locations.FLOORS for name in names]

