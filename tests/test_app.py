"""재물조사 배포판 API 테스트.

예시(더미) 자산대장·공간목록·빈 재물조사표로 실행한다. DB와 출력물은 임시 폴더로 격리한다.
"""
import io
import os
import sys
import tempfile
import unittest
from datetime import timedelta
from pathlib import Path
from unittest import mock

ROOT = Path(__file__).resolve().parents[1]
RUNTIME = ROOT / f".runtime-py{sys.version_info.major}{sys.version_info.minor}"
if RUNTIME.exists():
    sys.path.insert(0, str(RUNTIME))
sys.path.insert(0, str(ROOT))

ACCESS = "246813"
ADMIN = "135792"
os.environ["ACCESS_CODE"] = ACCESS
os.environ["ADMIN_CODE"] = ADMIN
os.environ["ORG_NAME"] = "테스트<복지관>"
os.environ["PASSKEY_RP_ID"] = ""
os.environ["PASSKEY_ORIGIN"] = ""

import openpyxl                            # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402
from app import clock, config, database, excel, locations, main  # noqa: E402
from app.catalog import catalog            # noqa: E402

DESK_7, DESK_8 = "01-01-001-0007", "01-01-001-0008"


class AppTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.db_patch = mock.patch.object(database, "DB_PATH", Path(self.temp.name) / "test.db")
        self.db_patch.start()
        main._FAILS.clear()
        catalog.load()
        database.init()
        self.client = TestClient(main.app)
        self.assertEqual(self.client.post("/api/session", data={"code": ACCESS}).status_code, 200)

    def tearDown(self):
        self.client.close()
        self.db_patch.stop()
        self.temp.cleanup()

    def _login_admin(self):
        self.assertEqual(self.client.post("/api/admin/session", data={"code": ADMIN}).status_code, 200)

    def _save(self, code, qty=1, status="활용", **extra):
        body = {"codes": [code], "primary_code": code, "actual_qty": qty, "status": status,
                "method": "list", **extra}
        return self.client.put("/api/results", json=body)

    # ── 자산대장 · 공간 ──
    def test_ledger_skips_subtotals_and_reads_ledger_columns(self):
        self.assertEqual(len(catalog.assets), 110)
        self.assertFalse(any(a["name"] in ("소계", "합계") for a in catalog.assets.values()))
        car = catalog.get("08-01-001")                 # 3마디 코드
        self.assertEqual(car["category"], "차량운반구")
        self.assertEqual(car["book_qty"], 1)
        self.assertEqual(car["space"], "주차장")
        self.assertNotIn("old_code", car)
        self.assertEqual(len(catalog.serial_group(DESK_7)), 8)
        self.assertEqual(catalog.serial_group("08-01-001"), [])

    def test_home_uses_space_list_and_collects_unlisted_locations(self):
        floors = self.client.get("/api/home").json()["floors"]
        self.assertEqual([f["floor"] for f in floors], ["1층", "2층", "3층", "별관", "외부", "기타"])
        third = {s["name"]: s for s in floors[2]["spaces"]}
        self.assertEqual(third["옥상정원"]["total"], 0)       # 목록 공간은 자산이 없어도 표시
        etc = {s["name"] for s in floors[-1]["spaces"]}
        self.assertEqual(etc, {"임시보관함", "미지정"})

    def test_spaces_fall_back_to_ledger_locations(self):
        with mock.patch.object(locations, "SOURCE_SPACES", Path(self.temp.name) / "none.xlsx"):
            catalog.load()
            groups = locations.official_groups()
        self.assertEqual(len(groups), 1)
        self.assertIn("임시보관함", groups[0]["spaces"])
        self.assertNotIn("미지정", groups[0]["spaces"])
        catalog.load()

    def test_search_and_lookup(self):
        self.assertGreater(self.client.get("/api/search", params={"q": "ㅅㅁㅇ"}).json()["total"], 0)
        self.assertGreater(self.client.get("/api/search", params={"q": "운영지원팀"}).json()["total"], 0)
        lookup = self.client.post("/api/lookup", json={"code": "01010010007"}).json()
        self.assertEqual(lookup["asset"]["code"], DESK_7)
        self.assertEqual(self.client.post("/api/lookup", json={"code": "7-1-1"}).json()["asset"]["code"], "07-01-001")

    # ── 저장 · 이력 ──
    def test_group_save_delete_and_audit_history(self):
        saved = self.client.put("/api/results", json={
            "codes": [DESK_7, DESK_8], "primary_code": DESK_8, "actual_qty": 2, "status": "활용",
            "confirmed_space": "로비", "method": "manual", "device": "홍길동 · iPhone",
        })
        self.assertEqual(saved.status_code, 200)
        self.assertEqual(database.get_result(DESK_7)["actual_qty"], 1)
        self.assertEqual(database.get_result(DESK_8)["device"], "홍길동 · iPhone")
        self.assertEqual(self.client.delete(f"/api/admin/results/{DESK_8}", params={"group": "true"}).status_code, 403)
        self._login_admin()
        self.assertEqual(self.client.delete(f"/api/admin/results/{DESK_8}", params={"group": "true"}).json()["count"], 2)
        self.assertTrue(any(h["action"] == "delete" for h in database.history_for(DESK_8)))

    def test_unlisted_location_requires_confirmed_space(self):
        self.assertEqual(self._save("07-02-001").status_code, 400)
        self.assertEqual(self._save("07-02-001", confirmed_space="3층창고").status_code, 200)

    def test_undo_and_conflict(self):
        first = self._save(DESK_7).json()
        stale = first["results"][0]["updated_at"]
        second = self._save(DESK_7, 2, expected_updated_at={DESK_7: stale})
        self.assertEqual(second.status_code, 200)
        self.assertEqual(self._save(DESK_7, 3, expected_updated_at={DESK_7: stale}).status_code, 409)
        undo = self.client.post("/api/results/undo", json={"token": second.json()["undo_token"]})
        self.assertEqual(undo.status_code, 200)
        self.assertEqual(database.get_result(DESK_7)["actual_qty"], 1)
        self.assertTrue(any(h["action"] == "restore" for h in database.history_for(DESK_7)))

    # ── 엑셀 출력 ──
    def _export(self, scope="all"):
        response = self.client.get("/api/admin/export", params={"scope": scope})
        self.assertEqual(response.status_code, 200, response.text)
        return openpyxl.load_workbook(io.BytesIO(response.content)).worksheets[0]

    def test_export_builds_rows_from_blank_template(self):
        self._save(DESK_7, 1, "폐품", reason="파손")
        self._login_admin()
        sheet = self._export()
        codes = [sheet.cell(r, 4).value for r in range(4, sheet.max_row + 1)]
        self.assertEqual(len(codes), 110)
        row = 4 + codes.index(DESK_7)
        values = [sheet.cell(row, c).value for c in range(1, 14)]
        self.assertEqual(values[0], clock.today_iso().replace("-", ""))
        self.assertEqual(values[1], "비품")
        self.assertEqual(values[7], 1)                      # 장부수
        self.assertEqual(values[8:13], [0, 1, 0, 0, "파손"])  # 사용가능·폐품·활용·과부족·사유
        other = 4 + codes.index(DESK_8)
        self.assertIsNone(sheet.cell(other, 1).value)        # 미실사 자산은 조사일자 비움
        only = self._export("surveyed")
        self.assertEqual([only.cell(r, 4).value for r in range(4, only.max_row + 1)], [DESK_7])

    def test_export_fills_prefilled_template_without_touching_other_columns(self):
        self._login_admin()
        self._save(DESK_8, 1, "사용가능")
        prefilled = Path(self.temp.name) / "filled.xlsx"
        wb = openpyxl.load_workbook(config.SOURCE_TEMPLATE)
        ws = wb.worksheets[0]
        ws.cell(4, 3).value, ws.cell(4, 4).value, ws.cell(4, 5).value = "FA0001", DESK_8, "기관 양식 자산명"
        wb.save(prefilled)
        with mock.patch.object(excel, "SOURCE_TEMPLATE", prefilled):
            sheet = self._export()
        self.assertEqual([sheet.cell(4, c).value for c in (3, 4, 5, 9)], ["FA0001", DESK_8, "기관 양식 자산명", 1])
        self.assertEqual(sheet.max_row, 4)

    # ── 인증 · 설정 ──
    def test_login_page_shows_org_name_and_bundled_font(self):
        anonymous = TestClient(main.app)
        page = anonymous.get("/").text
        self.assertIn("테스트&lt;복지관&gt;", page)
        self.assertNotIn("cdn.jsdelivr", page)
        self.assertEqual(anonymous.get("/fonts/PretendardVariable.woff2").status_code, 200)
        self.assertEqual(anonymous.get("/api/home").status_code, 401)
        index = self.client.get("/").text
        self.assertIn('data-org="테스트&lt;복지관&gt;"', index)

    def test_repeated_wrong_codes_lock_login(self):
        anonymous = TestClient(main.app)
        for _ in range(main.FAIL_LIMIT):
            self.assertEqual(anonymous.post("/api/session", data={"code": "000000"}).status_code, 401)
        self.assertEqual(anonymous.post("/api/session", data={"code": ACCESS}).status_code, 429)

    def test_code_rules(self):
        self.assertEqual(config.code_problems(ACCESS, ADMIN), [])
        self.assertEqual(len(config.code_problems("", ADMIN)), 1)
        self.assertEqual(len(config.code_problems("1234", ADMIN)), 1)
        self.assertEqual(len(config.code_problems(ACCESS, ACCESS)), 1)

    def test_logo_can_be_replaced(self):
        self.assertEqual(self.client.get("/brand/logo").headers["content-type"], "image/svg+xml")
        branding = Path(self.temp.name) / "branding"
        branding.mkdir()
        (branding / "logo.png").write_bytes(b"\x89PNG\r\n\x1a\n")
        with mock.patch.object(config, "BRANDING_DIR", branding):
            self.assertEqual(self.client.get("/brand/logo").headers["content-type"], "image/png")

    def test_passkey_disabled_without_address(self):
        self._login_admin()
        self.assertFalse(self.client.get("/api/admin/passkeys").json()["enabled"])
        self.assertEqual(self.client.post("/api/admin/passkeys/register/options").status_code, 404)
        self.assertFalse(self.client.get("/api/health").json()["passkey"])

    def test_database_newer_than_app_is_refused(self):
        with database.connect() as conn:
            conn.execute("UPDATE meta SET value=? WHERE key='schema_version'", (str(config.SCHEMA_VERSION + 1),))
        with self.assertRaises(RuntimeError):
            database.init()

    def test_dates_use_kst(self):
        self.assertEqual(clock.now().utcoffset(), timedelta(hours=9))
        saved = self._save(DESK_7).json()["results"][0]
        self.assertEqual(saved["survey_date"], clock.today_iso())


if __name__ == "__main__":
    unittest.main()
