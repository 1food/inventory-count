"""설정 화면(첫 실행 설정·기관 설정) 테스트. .env·자산 파일은 모두 임시 폴더로 바꿔 쓴다."""
import io
import os
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

ROOT = Path(__file__).resolve().parents[1]
RUNTIME = ROOT / f".runtime-py{sys.version_info.major}{sys.version_info.minor}"
if RUNTIME.exists():
    sys.path.insert(0, str(RUNTIME))
sys.path.insert(0, str(ROOT))
os.environ.setdefault("ACCESS_CODE", "246813")
os.environ.setdefault("ADMIN_CODE", "135792")

import openpyxl                            # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402
from app import config, database, main, setup  # noqa: E402
from app.catalog import catalog            # noqa: E402

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


class SetupTests(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        (self.tmp / "data").mkdir()
        assets = self.tmp / "assets"
        assets.mkdir()
        self.env = self.tmp / ".env"
        self.env.write_text("# 기관명 설명\nORG_NAME=\nACCESS_CODE=\nADMIN_CODE=\nOTHER=keep\n", encoding="utf-8")
        self.patches = [
            mock.patch.object(database, "DB_PATH", self.tmp / "data" / "survey.db"),
            mock.patch.object(setup, "ENV_PATH", self.env),
            mock.patch.object(setup, "SOURCE_SPACES", assets / "공간목록.xlsx"),
            mock.patch.object(setup, "PENDING", {
                "ledger": (assets / "l.upload.xlsx", assets / "고정자산대장.xlsx"),
                "template": (assets / "t.upload.xlsx", assets / "재물조사표.xlsx")}),
            mock.patch.object(setup, "BRANDING_DIR", self.tmp / "branding"),
            mock.patch.object(main, "BACKUP_DIR", self.tmp / "backups"),
            mock.patch.object(main, "DATA_DIR", self.tmp / "data"),
            mock.patch.object(main, "_schedule_restart"),
        ]
        for p in self.patches:
            p.start()
        self.restart = main._schedule_restart
        catalog.load()
        database.init()
        main.STATE.update(setup=True, problems=["코드 없음"])
        main._FAILS.clear()
        self.client = TestClient(main.app)

    def tearDown(self):
        main.STATE.update(setup=False, problems=[])
        self.client.close()
        for p in self.patches:
            p.stop()

    def local(self):
        return mock.patch.object(main, "_is_local", lambda request: True)

    def test_write_env_keeps_comments_and_other_keys(self):
        self.assertTrue(setup.write_env({"ORG_NAME": "가나\n복지관", "NEW": "1"}))
        text = self.env.read_text(encoding="utf-8")
        self.assertIn("# 기관명 설명", text)
        self.assertIn("ORG_NAME=가나복지관", text)   # 줄바꿈은 지운다
        self.assertIn("OTHER=keep", text)
        self.assertIn("NEW=1", text)
        self.assertFalse(setup.write_env({"ORG_NAME": "가나복지관"}))

    def test_setup_mode_only_opens_setup_page_on_server_pc(self):
        self.assertEqual(self.client.get("/").status_code, 503)             # 다른 기기
        self.assertEqual(self.client.get("/api/home").status_code, 503)
        self.assertEqual(self.client.get("/api/setup/state").status_code, 403)
        with self.local():
            self.assertEqual(self.client.get("/", follow_redirects=False).headers["location"], "/setup")
            state = self.client.get("/api/setup/state").json()
        self.assertTrue(state["setup_mode"])
        self.assertEqual(state["ledger"]["assets"], 110)
        self.assertTrue(state["ledger"]["sample"])

    def test_cloudflare_traffic_is_not_local(self):
        request = mock.Mock(client=mock.Mock(host="127.0.0.1"), headers={"cf-connecting-ip": "1.2.3.4"})
        self.assertFalse(main._is_local(request))
        request.headers = {}
        self.assertTrue(main._is_local(request))

    def test_ledger_upload_checks_before_saving(self):
        bad = io.BytesIO()
        workbook = openpyxl.Workbook()
        workbook.active.append(["이름", "수량"])
        workbook.save(bad)
        with self.local():
            rejected = self.client.post("/api/setup/ledger", files={"file": ("a.xlsx", bad.getvalue(), XLSX)})
            self.assertEqual(rejected.status_code, 400)
            self.assertIn("자산코드", rejected.json()["detail"])
            self.assertFalse(setup.PENDING["ledger"][0].exists())
            good = self.client.post("/api/setup/ledger", files={
                "file": ("b.xlsx", config.SOURCE_MASTER.read_bytes(), XLSX)}).json()
        self.assertEqual(good["assets"], 110)
        self.assertIn("주차장", [l["name"] for l in good["locations"]])
        self.assertTrue(setup.PENDING["ledger"][0].exists())
        self.assertEqual(good["diff"], {"added": 0, "removed": 0, "changed": 0,
                                        "examples": {"added": [], "removed": [], "changed": []}})

    def test_ledger_diff_counts_added_removed_changed(self):
        current = {"A-1": {"name": "책상", "raw_location": "로비", "book_qty": 1, "category": "비품"},
                   "A-2": {"name": "의자", "raw_location": "로비", "book_qty": 1, "category": "비품"}}
        rows = [{"code": "A-1", "name": "책상", "raw_location": "창고", "book_qty": 1, "category": "비품"},
                {"code": "A-3", "name": "모니터", "raw_location": "로비", "book_qty": 1, "category": "비품"}]
        with mock.patch.object(setup, "read_ledger", lambda path: rows):
            diff = setup.check_ledger(Path("x.xlsx"), current=current)["diff"]
        self.assertEqual((diff["added"], diff["removed"], diff["changed"]), (1, 1, 1))
        self.assertEqual(diff["examples"]["changed"], ["A-1 책상"])

    def test_save_validates_codes_then_writes_everything(self):
        body = {"org": "테스트복지관", "access_code": "1234", "admin_code": "135792",
                "spaces": [{"floor": "1층", "name": "로비"}, {"floor": "", "name": "창고"}]}
        with self.local():
            self.assertEqual(self.client.post("/api/setup/save", json=body).status_code, 400)
            self.client.post("/api/setup/ledger", files={"file": ("b.xlsx", config.SOURCE_MASTER.read_bytes(), XLSX)})
            body["access_code"] = "246813"
            saved = self.client.post("/api/setup/save", json=body).json()
        self.assertTrue(saved["restart"])
        self.restart.assert_called_once()
        text = self.env.read_text(encoding="utf-8")
        self.assertIn("ORG_NAME=테스트복지관", text)
        self.assertIn("ACCESS_CODE=246813", text)
        self.assertTrue(setup.PENDING["ledger"][1].exists())          # 올린 대장이 적용됨
        rows = list(openpyxl.load_workbook(setup.SOURCE_SPACES).active.iter_rows(values_only=True))
        self.assertEqual(rows[1:], [("1층", "로비", None), ("공간", "창고", None)])

    def test_after_setup_requires_admin_and_clear_backs_up_first(self):
        main.STATE.update(setup=False, problems=[])
        self.client.post("/api/session", data={"code": main.ACCESS_CODE})
        self.assertEqual(self.client.get("/api/setup/state").status_code, 403)
        self.client.post("/api/admin/session", data={"code": main.ADMIN_CODE})
        self.client.put("/api/results", json={"codes": ["08-01-001"], "primary_code": "08-01-001",
                                              "actual_qty": 1, "status": "활용", "method": "list"})
        self.assertEqual(self.client.post("/api/setup/clear-results", json={"confirm": "네"}).status_code, 400)
        cleared = self.client.post("/api/setup/clear-results", json={"confirm": "지우기"}).json()
        self.assertEqual(database.all_results(), {})
        self.assertTrue((Path(cleared["backup"]) / "survey.db").exists())
        download = self.client.get("/api/setup/backup/download")
        self.assertEqual(download.status_code, 200)
        self.assertTrue(download.content.startswith(b"PK"))
        self.assertIn("attachment", download.headers["content-disposition"])


if __name__ == "__main__":
    unittest.main()
